import os
import io
import re
import json
import tempfile
import datetime
import subprocess
import openpyxl
import pdfplumber
from flask import Flask, request, jsonify, send_file
from flask_cors import CORS

app = Flask(__name__)
CORS(app)

# テンプレートは「日報ほか/テンプレート/」を優先し、なければアプリ内templates/を使用
_PRIMARY_TEMPLATE = os.path.expanduser('~/Desktop/日報ほか/テンプレート/見積書サンプル.xlsx')
_FALLBACK_TEMPLATE = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'templates', '見積書サンプル.xlsx')
TEMPLATE_PATH = _PRIMARY_TEMPLATE if os.path.exists(_PRIMARY_TEMPLATE) else _FALLBACK_TEMPLATE
OCR_SCRIPT = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'ocr_pdf.swift')


# ─────────────────────────────────────────────
# OCR (Swift + Vision framework)
# ─────────────────────────────────────────────

def ocr_page_to_rows(pdf_path, page_index):
    """OCR a PDF page via Swift/Vision, return rows of items sorted left-to-right, top-to-bottom."""
    result = subprocess.run(
        ['swift', OCR_SCRIPT, pdf_path, str(page_index)],
        capture_output=True, text=True, timeout=90
    )
    if not result.stdout.strip():
        return []
    items = json.loads(result.stdout)
    # Sort top-to-bottom (Vision y=0 is bottom, so sort descending y)
    items.sort(key=lambda x: -x['y'])
    rows = []
    for item in items:
        placed = False
        for row in rows:
            if abs(row[0]['y'] - item['y']) < 0.018:
                row.append(item)
                placed = True
                break
        if not placed:
            rows.append([item])
    for row in rows:
        row.sort(key=lambda x: x['x'])
    return rows


def assign_to_columns(row_items, col_headers):
    """Map each cell to the nearest column header by x-position."""
    result = {}
    for item in row_items:
        best = min(col_headers, key=lambda h: abs(h['x'] - item['x']))
        name = best['text']
        if name in result:
            result[name] += ' ' + item['text']
        else:
            result[name] = item['text']
    return result


def normalize_col_name(text):
    """Normalize OCR-garbled column header names."""
    t = text.strip()
    # 商品コード (separate column) → keep as '商品コード'
    if re.search(r'商品[コゴ]ード', t) and '商品名' not in t:
        return '商品コード'
    maps = [
        ('商品コード商品名', '商品名'), ('商品名', '商品名'),
        ('入り数', '規格'), ('規格', '規格'), ('規楮', '規格'),
        ('荷姿', '荷姿'),
        ('営業売', '単価'), ('売価', '単価'),
        ('JANコード', 'JAN'), ('JAN', 'JAN'),
        ('賞味期限', '賞味期限'), ('賞味', '賞味期限'),
        ('備考', '備考'), ('儁考', '備考'),
    ]
    for key, val in maps:
        if key in t:
            return val
    return t


def clean_price(text):
    if not text:
        return None
    # Pattern with 円 suffix: "2,300円" or "k1,780円"
    m = re.search(r'k?([\d,]+)円', text)
    if m:
        return int(m.group(1).replace(',', ''))
    # Comma-grouped number without 円: "2,300" or "1,780"
    m = re.search(r'\b(\d{1,3}(?:,\d{3})+)\b', text)
    if m:
        return int(m.group(1).replace(',', ''))
    # Last standalone integer
    nums = re.findall(r'\d+', text)
    if nums:
        try:
            return int(nums[-1])
        except ValueError:
            pass
    return None


def clean_jan(text):
    if not text:
        return ''
    s = text.strip().replace(' ', '')
    if re.match(r'^[ー\-－―一—]+$', s):
        return ''
    # Keep only digits
    digits = re.sub(r'[^\d]', '', s)
    if len(digits) >= 8:
        return digits
    return ''


def is_product_code(text):
    """Check if text starts with a 6-digit product code."""
    return bool(re.match(r'^\d{6}', text.strip()))


def extract_name_from_cell(text):
    """Strip leading product code from combined code+name cell."""
    return re.sub(r'^\d{6}[\s:：．.]*', '', text).strip()


SKIP_PATTERNS = [
    '工場出荷目標', '出荷目標', '諸条件', '改訂版', '【拡販商品】',
    '甲単位', '1/1', '2026/3', '07.04',
]

def should_skip_row(combined):
    if any(kw in combined for kw in SKIP_PATTERNS):
        return True
    # Footer with digit/date patterns
    if re.match(r'^\d{2}[./]\d{2}[./]\d{2}', combined):
        return True
    # Lines like "さっぱりやみつき小魚5甲" (product + 甲)
    if re.search(r'[^\d]\d+甲$', combined):
        return True
    # Very short noise lines (single kanji/symbol rows)
    if len(combined.strip()) <= 2:
        return True
    return False


def clean_haishi(text):
    """Normalize 荷姿 and strip any price that leaked in."""
    if not text:
        return ''
    # Normalize spacing around ×
    h = re.sub(r'(\d+)\s*[xXｘ×]\s*(\d+)', r'\1×\2', text)
    # Strip trailing price like " 270円" or " k1,780円"
    h = re.sub(r'\s+k?[\d,]+円\s*$', '', h).strip()
    # Strip standalone price-like fragments
    h = re.sub(r'^k?[\d,]+円$', '', h).strip()
    return h


def parse_ocr_rows(rows):
    """Parse reconstructed OCR rows into product list."""
    # Find header row: must contain '商品名' and ('荷姿' or '営業売')
    header_idx = None
    for i, row in enumerate(rows):
        combined = ' '.join(item['text'] for item in row)
        if '商品名' in combined and ('荷姿' in combined or '営業売' in combined):
            header_idx = i
            break
    if header_idx is None:
        return []

    # Build normalized column headers with their x positions
    raw_headers = rows[header_idx]
    col_headers = []
    seen_names = set()
    for item in raw_headers:
        norm = normalize_col_name(item['text'])
        if norm not in seen_names:
            seen_names.add(norm)
            col_headers.append({'text': norm, 'x': item['x']})

    # x-position of 荷姿 and 単価 columns for range-based extraction
    haishi_hdr = next((h for h in col_headers if h['text'] == '荷姿'), None)
    kakaku_hdr = next((h for h in col_headers if h['text'] == '単価'), None)

    products = []
    current = None

    for row in rows[header_idx + 1:]:
        if not row:
            continue
        combined = ' '.join(item['text'] for item in row)
        if should_skip_row(combined):
            continue

        mapped = assign_to_columns(row, col_headers)

        # ── Detect new product row ──
        has_separate_code_col = '商品コード' in seen_names
        if has_separate_code_col:
            # 五色浜 / 営業部仕入課 format: 商品コード and 商品名 are separate columns
            code_raw = mapped.get('商品コード', '')
            is_new = is_product_code(code_raw)
            name_raw = mapped.get('商品名', '')
        else:
            # セトクイーン format: code+name merged in one column
            name_raw = mapped.get('商品名', '')
            is_new = bool(name_raw) and is_product_code(name_raw)

        # ── Continuation row ──
        if not is_new:
            if current:
                if not current['荷姿'] and mapped.get('荷姿'):
                    current['荷姿'] = clean_haishi(mapped['荷姿'])
                if not current['規格'] and mapped.get('規格'):
                    current['規格'] = mapped['規格']
                biko = mapped.get('備考', '').strip()
                if biko and biko not in current['備考']:
                    current['備考'] += (' ' + biko if current['備考'] else biko)
            continue

        # ── New product row ──
        if current:
            products.append(current)

        if has_separate_code_col:
            name = name_raw.strip()
        else:
            name = extract_name_from_cell(name_raw)
        name = re.sub(r'【.*?】', '', name).strip()
        if not name:
            current = None
            continue

        # 荷姿: collect cells strictly between 荷姿-col and 単価-col x-ranges
        haishi = ''
        if haishi_hdr and kakaku_hdr:
            lo = haishi_hdr['x'] - 0.04
            hi = kakaku_hdr['x'] - 0.01
            parts = [item['text'] for item in row if lo <= item['x'] <= hi]
            haishi = clean_haishi(' '.join(parts))
        if not haishi:
            haishi = clean_haishi(mapped.get('荷姿', ''))

        current = {
            '商品名': name,
            '規格': mapped.get('規格', ''),
            '荷姿': haishi,
            '単価': clean_price(mapped.get('単価', '')),
            'JANコード': clean_jan(mapped.get('JAN', '')),
            '賞味期限': re.sub(r'[\s　]', '', mapped.get('賞味期限', ''))
                          .replace('E', '日').replace('B', '日').replace('F', '日'),
            '備考': mapped.get('備考', '').strip(),
        }

    if current:
        products.append(current)

    return products


# ─────────────────────────────────────────────
# PDF extraction dispatcher
# ─────────────────────────────────────────────

def is_scanned_page(page):
    """Check if a pdfplumber page has no extractable text (scanned image)."""
    text = page.extract_text()
    return not text or not text.strip()


def extract_products_from_pdf(pdf_path):
    """Extract products from all relevant pages of the PDF."""
    products = []

    with pdfplumber.open(pdf_path) as pdf:
        for page_idx, page in enumerate(pdf.pages):
            # Check if page has extractable text
            if is_scanned_page(page):
                # OCR via Swift/Vision
                try:
                    rows = ocr_page_to_rows(pdf_path, page_idx)
                    page_products = parse_ocr_rows(rows)
                    products.extend(page_products)
                except Exception:
                    pass

    # Deduplicate by 商品名
    seen = set()
    unique = []
    for p in products:
        key = p['商品名']
        if key not in seen:
            seen.add(key)
            unique.append(p)

    return unique


# ─────────────────────────────────────────────
# Excel output
# ─────────────────────────────────────────────

def write_products_to_sheet(ws, batch, line_number_start):
    for i, product in enumerate(batch):
        row = 13 + i
        ws.cell(row=row, column=2).value = line_number_start + i
        ws.cell(row=row, column=3).value = product['商品名']
        ws.cell(row=row, column=10).value = product['規格']
        ws.cell(row=row, column=12).value = product['荷姿']
        ws.cell(row=row, column=15).value = '＠'
        price = product.get('単価')
        ws.cell(row=row, column=16).value = int(price) if price else ''
        ws.cell(row=row, column=18).value = product['賞味期限']
        jan = product.get('JANコード', '')
        ws.cell(row=row, column=20).value = str(jan) if jan else ''
        ws.cell(row=row, column=24).value = product.get('備考', '')


def clear_data_rows(ws):
    for row in range(13, 23):
        for col in [2, 3, 10, 12, 15, 16, 18, 20, 24]:
            ws.cell(row=row, column=col).value = ''


def build_excel(products):
    wb = openpyxl.load_workbook(TEMPLATE_PATH, keep_vba=False)
    ws1 = wb.worksheets[0]
    batches = [products[i:i+10] for i in range(0, len(products), 10)]
    if not batches:
        batches = [[]]

    write_products_to_sheet(ws1, batches[0], 1)
    for idx, batch in enumerate(batches[1:], start=1):
        new_ws = wb.copy_worksheet(ws1)
        new_ws.title = f'Sheet{idx + 1}'
        clear_data_rows(new_ws)
        write_products_to_sheet(new_ws, batch, idx * 10 + 1)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ─────────────────────────────────────────────
# Routes
# ─────────────────────────────────────────────

@app.route('/')
def index():
    try:
        with open(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'index.html'), encoding='utf-8') as f:
            return f.read(), 200, {'Content-Type': 'text/html; charset=utf-8'}
    except FileNotFoundError:
        return 'index.html not found', 404


@app.route('/upload', methods=['POST'])
def upload():
    if 'pdf' not in request.files:
        return jsonify({'error': 'PDFファイルが見つかりません'}), 400
    pdf_file = request.files['pdf']
    if not pdf_file.filename:
        return jsonify({'error': 'ファイルが選択されていません'}), 400
    if not pdf_file.filename.lower().endswith('.pdf'):
        return jsonify({'error': 'PDFファイルのみ対応しています'}), 400

    tmp_path = None
    try:
        with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as tmp:
            pdf_file.save(tmp.name)
            tmp_path = tmp.name

        try:
            products = extract_products_from_pdf(tmp_path)
        except Exception as e:
            return jsonify({'error': f'PDF読み取りエラー: {str(e)}'}), 500

        if not products:
            return jsonify({'error': 'PDFから商品データを抽出できませんでした。'}), 400

        try:
            excel_buf = build_excel(products)
        except Exception as e:
            return jsonify({'error': f'Excel生成エラー: {str(e)}'}), 500

        today = datetime.date.today().strftime('%Y%m%d')
        filename = f'見積書_完成版_{today}.xlsx'
        return send_file(
            excel_buf,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=filename
        )
    except Exception as e:
        return jsonify({'error': f'予期しないエラー: {str(e)}'}), 500
    finally:
        if tmp_path and os.path.exists(tmp_path):
            try:
                os.unlink(tmp_path)
            except Exception:
                pass


if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5001)
